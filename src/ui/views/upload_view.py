"""
upload_view.py — Production-grade upload flow with validated text extraction.

FIXES APPLIED:
  1. extract_text(): replaced bare `except:` with typed exception handling + logging.
  2. Added empty-text validation after extraction — warns user instead of silently continuing.
  3. "AI Ready" status is now only set after intelligence extraction actually succeeds
     (summary is not the failure sentinel string).
  4. Visible error messages shown for extraction or AI failures instead of silent fallbacks.
  5. Vector store building is now triggered and stored in session_state for use by chat.
"""
import io
import time
import uuid
import logging
import traceback

import streamlit as st

from constants import FOREST, SAND, RUST, SAGE, CREAM, WHITE, ALLOWED_FILE_TYPES

logger = logging.getLogger("upload_view")


# ── Text Extraction ────────────────────────────────────────────────────────────

def extract_text(uploaded_file) -> tuple[str, str | None]:
    """
    Extract plain text from an uploaded file.

    Returns:
        (text, error_msg) — error_msg is None on success.
    """
    ext  = uploaded_file.name.rsplit(".", 1)[-1].lower()
    data = uploaded_file.read()

    if ext == "pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            pages  = [page.extract_text() or "" for page in reader.pages]
            text   = "\n".join(pages).strip()
            if not text:
                return "", (
                    "PDF appears to be image-based (scanned). "
                    "Only text-based PDFs are supported for AI analysis."
                )
            logger.info(f"PDF extracted: {len(reader.pages)} pages, {len(text)} chars")
            return text, None
        except ImportError:
            logger.error("pypdf is not installed.")
            return "", "The 'pypdf' library is missing. Please run 'pip install pypdf'."
        except Exception:
            logger.error(f"PDF extraction failed:\n{traceback.format_exc()}")
            return "", "PDF could not be parsed. It may be corrupted or password-protected."

    elif ext == "docx":
        try:
            import docx
            doc  = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            if not text:
                return "", "DOCX file appears to contain no extractable text."
            logger.info(f"DOCX extracted: {len(text)} chars")
            return text, None
        except ImportError:
            logger.error("python-docx is not installed.")
            return "", "The 'python-docx' library is missing. Please run 'pip install python-docx'."
        except Exception:
            logger.error(f"DOCX extraction failed:\n{traceback.format_exc()}")
            return "", "DOCX could not be parsed. It may be corrupted."

    elif ext in ("txt", "md"):
        try:
            text = data.decode("utf-8", errors="replace").strip()
            if not text:
                return "", "Text file is empty."
            return text, None
        except Exception:
            logger.error(f"Text file decode failed:\n{traceback.format_exc()}")
            return "", "Text file could not be decoded."

    else:
        return "", f"Unsupported file type: .{ext}"


# ── Build Vector Store ─────────────────────────────────────────────────────────

def _build_vector_store(text: str, uploaded_file_name: str):
    """
    Chunk the text and build a FAISS vector store.
    Returns the vector store on success, None on failure.
    """
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        from langchain_core.documents import Document
        from backend.embeddings import get_vector_store

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=500, chunk_overlap=50,
            length_function=len, is_separator_regex=False
        )
        chunks = splitter.split_text(text)
        if not chunks:
            logger.warning("Text produced 0 chunks — skipping vector store build.")
            return None

        docs = [
            Document(
                page_content=chunk.strip(),
                metadata={"source": uploaded_file_name, "chunk_index": i}
            )
            for i, chunk in enumerate(chunks)
        ]
        logger.info(f"Building vector store: {len(docs)} chunks")
        vs = get_vector_store(docs)
        logger.info("Vector store built successfully.")
        return vs

    except Exception:
        logger.error(f"Vector store build failed:\n{traceback.format_exc()}")
        return None


# ── Render ─────────────────────────────────────────────────────────────────────

def render():
    st.markdown("<div style='height: 4vh'></div>", unsafe_allow_html=True)

    col1, col2, col3 = st.columns([1, 4, 1])

    with col2:
        st.markdown(f"""
        <div style="text-align: center; margin-bottom: 32px;">
            <div style="width: 56px; height: 56px; background: linear-gradient(135deg, {SAGE}, {FOREST}); border-radius: 16px; margin: 0 auto 16px; display: flex; align-items: center; justify-content: center; font-size: 28px; box-shadow: 0 8px 16px rgba(0,0,0,0.1);">🌿</div>
            <h1 style="font-family: 'Playfair Display', serif; color: {FOREST}; font-size: 42px; font-weight: 700; margin-bottom: 8px;">Prepo AI</h1>
            <p style="color: {SAGE}; font-size: 16px; max-width: 400px; margin: 0 auto; line-height: 1.5;">The AI understands your material and builds an intelligent workspace around it.</p>
        </div>
        """, unsafe_allow_html=True)

        st.markdown(f"""
        <div style="background: {WHITE}; border-radius: 16px; padding: 32px; box-shadow: 0 4px 24px rgba(0,0,0,0.04); border: 1px solid {SAND}30;">
        """, unsafe_allow_html=True)

        uploaded_file = st.file_uploader(
            "Upload your study material",
            type=ALLOWED_FILE_TYPES,
            help=f"Supported: {', '.join(ALLOWED_FILE_TYPES).upper()}",
            label_visibility="collapsed"
        )

        if not uploaded_file:
            st.markdown(f"""
            <div style="text-align: center; margin-top: 24px; color: {SAGE}; font-size: 13.5px;">
                <div style="font-weight: 600; margin-bottom: 12px; color: {FOREST};">💡 Try uploading:</div>
                <span style="background: {CREAM}; padding: 6px 12px; border-radius: 20px; border: 1px solid {SAND}40; margin: 4px; display: inline-block;">Lecture Notes</span>
                <span style="background: {CREAM}; padding: 6px 12px; border-radius: 20px; border: 1px solid {SAND}40; margin: 4px; display: inline-block;">Textbook Chapters</span>
                <span style="background: {CREAM}; padding: 6px 12px; border-radius: 20px; border: 1px solid {SAND}40; margin: 4px; display: inline-block;">Research Papers</span>
            </div>
            """, unsafe_allow_html=True)

        st.markdown("</div>", unsafe_allow_html=True)

        if uploaded_file:
            st.markdown("<div style='height: 24px'></div>", unsafe_allow_html=True)

            with st.status("Building your intelligent workspace...", expanded=True) as status:

                # ── Step 1: Extract Text ──────────────────────────────────────
                st.write("📄 Reading document...")
                text_content, extract_error = extract_text(uploaded_file)

                if extract_error:
                    status.update(label="Document Error", state="error", expanded=True)
                    st.error(f"❌ {extract_error}")
                    st.stop()

                st.write(f"✅ Extracted {len(text_content):,} characters from document.")
                time.sleep(0.3)

                # ── Step 2: AI Intelligence ───────────────────────────────────
                st.write("🧠 Extracting core intelligence...")
                from backend.intelligence import extract_document_intelligence
                intelligence = extract_document_intelligence(text_content)

                ai_error = intelligence.get("error")
                ai_ok = (
                    ai_error is None
                    and intelligence.get("summary", "").strip()
                    and intelligence["summary"] not in {
                        "Could not generate summary — the AI service returned an error. Check logs for details.",
                        "No text could be extracted from this document. It may be a scanned image-based PDF.",
                    }
                )

                if ai_ok:
                    st.write("✅ AI summary and topics generated.")
                else:
                    st.warning(f"⚠️ AI summary could not be generated: {ai_error or 'Unknown error'}. You can still use the document.")

                time.sleep(0.3)

                # ── Step 3: Build Vector Store ────────────────────────────────
                st.write("🔍 Building vector search index...")
                vector_store = _build_vector_store(text_content, uploaded_file.name)

                if vector_store:
                    st.write(f"✅ Vector search index ready.")
                else:
                    st.warning("⚠️ Vector search index could not be built. Contextual chat will use raw text instead.")

                time.sleep(0.3)

                # ── Step 4: Save to Session State ─────────────────────────────
                st.write("✨ Preparing AI workspace...")
                doc_id = str(uuid.uuid4())

                if "uploaded_docs" not in st.session_state:
                    st.session_state.uploaded_docs = {}

                st.session_state.uploaded_docs[doc_id] = {
                    "name":          uploaded_file.name,
                    "text":          text_content,
                    "summary":       intelligence.get("summary", "Summary not available."),
                    "topics":        intelligence.get("topics",  []),
                    "upload_time":   time.time(),
                    "ai_ready":      ai_ok,
                    "vector_store":  vector_store,       # stored for contextual chat
                    "char_count":    len(text_content),
                }

                st.session_state.active_document_id = doc_id
                st.session_state.app_mode           = "workspace"
                st.session_state.workspace_view     = "overview"

                status.update(
                    label="Workspace Ready!" if ai_ok else "Workspace Ready (AI summary unavailable)",
                    state="complete",
                    expanded=False
                )

            st.rerun()